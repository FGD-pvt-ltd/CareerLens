import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Reference PPT Color Palette
    BG_COLOR = RGBColor(250, 249, 245)         # Warm Cream / Off-white (#FAF9F5)
    CARD_BG = RGBColor(255, 255, 255)          # Pure White (#FFFFFF)
    CARD_BORDER = RGBColor(226, 232, 240)      # Slate Border (#E2E8F0)
    CARD_BORDER_DARK = RGBColor(203, 213, 225) # Darker Slate (#CBD5E1)
    
    TEXT_MAIN = RGBColor(17, 24, 39)           # Deep Charcoal / Obsidian (#111827)
    TEXT_MUTED = RGBColor(75, 85, 99)          # Cool Slate Gray (#4B5563)
    TEXT_LIGHT = RGBColor(107, 114, 128)       # Light Slate (#6B7280)
    
    ACCENT_BLUE = RGBColor(37, 99, 235)        # Electric Blue (#2563EB)
    ACCENT_BLUE_BG = RGBColor(239, 246, 255)    # Light Blue Tint (#EFF6FF)
    
    ACCENT_ROSE = RGBColor(225, 29, 72)        # Rose / Hot Pink (#E11D48)
    ACCENT_ROSE_BG = RGBColor(255, 241, 242)    # Light Rose Tint (#FFF1F2)
    
    ACCENT_AMBER = RGBColor(217, 119, 6)       # Amber / Orange (#D97706)
    ACCENT_AMBER_BG = RGBColor(254, 243, 199)   # Light Amber Tint (#FEF3C7)
    
    ACCENT_EMERALD = RGBColor(16, 185, 129)    # Emerald Green (#10B981)
    ACCENT_EMERALD_BG = RGBColor(236, 253, 245) # Light Emerald (#ECFDF5)
    
    FONT_FAMILY = "Segoe UI"
    
    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_COLOR
        
        # Subtle bottom footer bar
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.1), Inches(11.733), Inches(0.3))
        tf = footer_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = "ProfiQ · AI-Powered Employability & Career Readiness Analyzer  |  DataQuest 3.0 · Round 1"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(9)
        p.font.color.rgb = TEXT_LIGHT

    def add_header(slide, badge_text, title_text, subtitle_text=""):
        # Category Pill / Badge
        badge_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.42), Inches(2.8), Inches(0.32))
        badge_shape.fill.solid()
        badge_shape.fill.fore_color.rgb = ACCENT_BLUE_BG
        badge_shape.line.color.rgb = RGBColor(191, 219, 254)
        badge_shape.line.width = Pt(1)
        
        btf = badge_shape.text_frame
        btf.vertical_anchor = MSO_ANCHOR.TOP
        btf.margin_left = btf.margin_right = 0
        btf.margin_top = Inches(0.04)
        btf.margin_bottom = 0
        bp = btf.paragraphs[0]
        bp.text = badge_text.upper()
        bp.alignment = PP_ALIGN.CENTER
        bp.font.name = FONT_FAMILY
        bp.font.size = Pt(9)
        bp.font.bold = True
        bp.font.color.rgb = ACCENT_BLUE
        
        # Title
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.733), Inches(0.85))
        tf = tb.text_frame
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = FONT_FAMILY
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN
        
        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.name = FONT_FAMILY
            p2.font.size = Pt(11)
            p2.font.color.rgb = TEXT_MUTED
            p2.space_before = Pt(3)

    # -------------------------------------------------------------
    # SLIDE 1: TITLE & COVER
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    
    # Event Pill
    ev_pill = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.55), Inches(3.2), Inches(0.34))
    ev_pill.fill.solid()
    ev_pill.fill.fore_color.rgb = ACCENT_ROSE_BG
    ev_pill.line.color.rgb = RGBColor(254, 205, 211)
    ev_tf = ev_pill.text_frame
    ev_tf.vertical_anchor = MSO_ANCHOR.TOP
    ev_tf.margin_left = ev_tf.margin_right = 0
    ev_tf.margin_top = Inches(0.04)
    ev_tf.margin_bottom = 0
    ev_p = ev_tf.paragraphs[0]
    ev_p.text = "DATAQUEST 3.0 · TRACK DQWL"
    ev_p.alignment = PP_ALIGN.CENTER
    ev_p.font.name = FONT_FAMILY
    ev_p.font.size = Pt(9.5)
    ev_p.font.bold = True
    ev_p.font.color.rgb = ACCENT_ROSE

    # Hero Title Box
    title_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.05), Inches(11.733), Inches(1.4))
    tf1 = title_box.text_frame
    tf1.vertical_anchor = MSO_ANCHOR.TOP
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0
    
    p = tf1.paragraphs[0]
    p.text = "ProfiQ"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(46)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN
    
    p_sub = tf1.add_paragraph()
    p_sub.text = "AI-Powered Employability & Career Readiness Analyzer"
    p_sub.font.name = FONT_FAMILY
    p_sub.font.size = Pt(17)
    p_sub.font.bold = True
    p_sub.font.color.rgb = ACCENT_BLUE
    p_sub.space_before = Pt(4)
    
    p_pitch = tf1.add_paragraph()
    p_pitch.text = "From Unverified Resume Claims to Observable Proof of Work — Bridging the Gap Between Student Skills and Industry Benchmarks"
    p_pitch.font.name = FONT_FAMILY
    p_pitch.font.size = Pt(12)
    p_pitch.font.color.rgb = TEXT_MUTED
    p_pitch.space_before = Pt(3)

    # Left Card: Core Project Capabilities
    card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.75), Inches(6.8), Inches(4.15))
    card1.fill.solid()
    card1.fill.fore_color.rgb = CARD_BG
    card1.line.color.rgb = CARD_BORDER
    card1.line.width = Pt(1)
    
    ctf1 = card1.text_frame
    ctf1.vertical_anchor = MSO_ANCHOR.TOP
    ctf1.word_wrap = True
    ctf1.margin_left = Inches(0.35)
    ctf1.margin_right = Inches(0.35)
    ctf1.margin_top = Inches(0.28)
    ctf1.margin_bottom = Inches(0.2)
    
    cp1 = ctf1.paragraphs[0]
    cp1.text = "CORE PROJECT HIGHLIGHTS"
    cp1.font.name = FONT_FAMILY
    cp1.font.size = Pt(12)
    cp1.font.bold = True
    cp1.font.color.rgb = ACCENT_BLUE
    
    bullets_s1 = [
        ("Evidence-Based Skill Verification: ", "Cross-validates claimed technologies against live GitHub repositories, commit velocity, and code architecture rather than passive resume keywords."),
        ("Deterministic Readiness Scoring: ", "Mathematically balances skill alignment (60%) and evidence authenticity (40%) to eliminate resume inflation and gaming."),
        ("Actionable Career Remediation: ", "Diagnoses candidate role deficits and generates automated, week-by-week milestone roadmaps with curated technical resources."),
        ("Institutional Placement Insights: ", "Provides university placement cells with batch-level analytics, cohort readiness distributions, and skill gap heatmaps to run targeted campus training.")
    ]
    
    for title, desc in bullets_s1:
        p_b = ctf1.add_paragraph()
        p_b.space_before = Pt(10)
        run1 = p_b.add_run()
        run1.text = "• " + title
        run1.font.name = FONT_FAMILY
        run1.font.size = Pt(10.5)
        run1.font.bold = True
        run1.font.color.rgb = TEXT_MAIN
        
        run2 = p_b.add_run()
        run2.text = desc
        run2.font.name = FONT_FAMILY
        run2.font.size = Pt(10)
        run2.font.color.rgb = TEXT_MUTED

    # Right Card: Team & Member Roles
    card2 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.85), Inches(2.75), Inches(4.68), Inches(4.15))
    card2.fill.solid()
    card2.fill.fore_color.rgb = CARD_BG
    card2.line.color.rgb = CARD_BORDER
    card2.line.width = Pt(1)
    
    ctf2 = card2.text_frame
    ctf2.vertical_anchor = MSO_ANCHOR.TOP
    ctf2.word_wrap = True
    ctf2.margin_left = Inches(0.35)
    ctf2.margin_right = Inches(0.35)
    ctf2.margin_top = Inches(0.28)
    ctf2.margin_bottom = Inches(0.2)
    
    cp2 = ctf2.paragraphs[0]
    cp2.text = "TEAM: [ 3Cuties ]"
    cp2.font.name = FONT_FAMILY
    cp2.font.size = Pt(12)
    cp2.font.bold = True
    cp2.font.color.rgb = ACCENT_AMBER
    
    members = [
        ("Aman Tyagi", "25BCE1710", "AI & ML Lead", "Google Gemini prompt orchestration, evidence verification heuristics, schema validation & scoring engines."),
        ("Siddhant Asati", "25BCE1633", "Backend & Systems Lead", "Node/Express REST APIs, MongoDB data layer, GitHub telemetry extraction & service normalization."),
        ("Sahaj Kumar", "25BCE1620", "Frontend & UX Lead", "React 18 + Vite interactive studio, real-time readiness dials, 3D visualizers & placement dashboard.")
    ]
    
    for name, reg, role, resp in members:
        p_m = ctf2.add_paragraph()
        p_m.space_before = Pt(12)
        
        r_name = p_m.add_run()
        r_name.text = f"{name} "
        r_name.font.name = FONT_FAMILY
        r_name.font.size = Pt(11)
        r_name.font.bold = True
        r_name.font.color.rgb = TEXT_MAIN
        
        r_reg = p_m.add_run()
        r_reg.text = f"({reg})\n"
        r_reg.font.name = FONT_FAMILY
        r_reg.font.size = Pt(10)
        r_reg.font.color.rgb = TEXT_LIGHT
        
        r_role = p_m.add_run()
        r_role.text = f"{role} · "
        r_role.font.name = FONT_FAMILY
        r_role.font.size = Pt(10)
        r_role.font.bold = True
        r_role.font.color.rgb = ACCENT_BLUE
        
        r_desc = p_m.add_run()
        r_desc.text = resp
        r_desc.font.name = FONT_FAMILY
        r_desc.font.size = Pt(9.5)
        r_desc.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 2: PROBLEM DEFINITION
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "01 · Problem Space", "The Employability Disconnect: Claims vs. Proof", 
               "Why traditional resumes, ATS keyword matchers, and unverified profiles fail modern engineering hiring.")

    cols_s2 = [
        ("01 · THE STUDENT DILEMMA", ACCENT_ROSE, [
            ("The Keyword Stuffing Trap: ", "Students list 20+ technologies on static resumes to bypass ATS filters, masking their true technical depth."),
            ("Fragmented Identity: ", "Proof of work is scattered across GitHub, LeetCode, portfolios, and courses with zero unified evaluation."),
            ("Reactive Rejections: ", "Students discover critical skill deficits only after failing technical placement rounds when it is too late to fix.")
        ]),
        ("02 · THE PLACEMENT BOTTLENECK", ACCENT_AMBER, [
            ("Batch Review Overload: ", "Placement cells must evaluate 1,000+ graduating students, relying on high-level GPA rather than code readiness."),
            ("Zero Cohort Telemetry: ", "Institutions lack aggregated data to know which frameworks or technologies students are deficient in."),
            ("Mismatched Drives: ", "Students are routed to unsuitable hiring partners, resulting in poor placement conversion rates.")
        ]),
        ("03 · THE RECRUITER FRICTION", ACCENT_BLUE, [
            ("Low Signal-to-Noise: ", "70%+ of applicants with polished resumes fail basic architectural and practical coding evaluations."),
            ("Manual Verification Burden: ", "Engineering managers spend hours auditing GitHub repos, commit histories, and code quality manually."),
            ("Expensive Mis-Hires: ", "Relying on self-reported resume claims creates inflated interview funnels and high recruitment rework.")
        ])
    ]

    card_w = Inches(3.68)
    card_h = Inches(4.15)
    card_top = Inches(1.8)

    for i, (col_title, accent_col, bullet_items) in enumerate(cols_s2):
        left_pos = Inches(0.8 + i * 4.02)
        c_shape = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, card_top, card_w, card_h)
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG
        c_shape.line.color.rgb = CARD_BORDER
        c_shape.line.width = Pt(1)
        
        ctf = c_shape.text_frame
        ctf.vertical_anchor = MSO_ANCHOR.TOP
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.28)
        ctf.margin_top = Inches(0.25)
        ctf.margin_bottom = Inches(0.2)
        
        cp = ctf.paragraphs[0]
        cp.text = col_title
        cp.font.name = FONT_FAMILY
        cp.font.size = Pt(11)
        cp.font.bold = True
        cp.font.color.rgb = accent_col
        
        for b_title, b_desc in bullet_items:
            bp = ctf.add_paragraph()
            bp.space_before = Pt(12)
            
            r1 = bp.add_run()
            r1.text = "• " + b_title
            r1.font.name = FONT_FAMILY
            r1.font.size = Pt(10.5)
            r1.font.bold = True
            r1.font.color.rgb = TEXT_MAIN
            
            r2 = bp.add_run()
            r2.text = b_desc
            r2.font.name = FONT_FAMILY
            r2.font.size = Pt(10)
            r2.font.color.rgb = TEXT_MUTED

    # Bottom Callout Banner
    bot_banner = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.12), Inches(11.733), Inches(0.82))
    bot_banner.fill.solid()
    bot_banner.fill.fore_color.rgb = RGBColor(241, 245, 249)
    bot_banner.line.color.rgb = CARD_BORDER_DARK
    bot_banner.line.width = Pt(1)
    
    btf = bot_banner.text_frame
    btf.vertical_anchor = MSO_ANCHOR.TOP
    btf.word_wrap = True
    btf.margin_left = Inches(0.3)
    btf.margin_right = Inches(0.3)
    btf.margin_top = Inches(0.12)
    btf.margin_bottom = Inches(0.1)
    
    bp_b = btf.paragraphs[0]
    bp_b.text = "THE ROOT CAUSE: \"A resume is an unverified claim. A GitHub repository is raw data.\""
    bp_b.font.name = FONT_FAMILY
    bp_b.font.size = Pt(11)
    bp_b.font.bold = True
    bp_b.font.color.rgb = ACCENT_BLUE
    
    bp_b2 = btf.add_paragraph()
    bp_b2.text = "No existing platform automatically cross-references self-reported skills against demonstrable code proof to generate an explainable, benchmarked readiness score with actionable remediation."
    bp_b2.font.name = FONT_FAMILY
    bp_b2.font.size = Pt(9.5)
    bp_b2.font.color.rgb = TEXT_MUTED
    bp_b2.space_before = Pt(3)

    # -------------------------------------------------------------
    # SLIDE 3: PROPOSED SOLUTION
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "02 · Proposed Solution", "ProfiQ: Evidence-Based Employability Intelligence", 
               "An end-to-end AI verification engine that converts raw developer claims into verifiable career readiness.")

    pipeline_steps = [
        ("1. INGEST & AGGREGATE", ACCENT_BLUE, [
            ("Resume Ingestion: ", "PDF parsing extracts candidate bio, education, claimed skills, and projects."),
            ("Developer Telemetry: ", "Fetches live GitHub repositories, language distributions, and commit frequency."),
            ("Target Role Benchmark: ", "Candidate selects intended role (e.g., Frontend Engineer, Full Stack, ML Engineer).")
        ]),
        ("2. AI VERIFY & BENCHMARK", ACCENT_ROSE, [
            ("Artifact Cross-Check: ", "Google Gemini cross-references claimed skills against actual repo code & dependencies."),
            ("Confidence Scoring: ", "Separates verified proof from unbacked claims using deterministic heuristics."),
            ("Role Matrix Comparison: ", "Evaluates depth against standardized critical, recommended, and optional skills.")
        ]),
        ("3. QUANTIFY & REMEDIATE", ACCENT_EMERALD, [
            ("Readiness Score (0-100): ", "Transparent, explainable score combining skill fit (60%) and proof depth (40%)."),
            ("Gap Identification: ", "Pinpoints exact missing technical dependencies required for target roles."),
            ("Milestone Roadmap: ", "Generates structured 4-week learning sprints with curated project objectives.")
        ])
    ]

    p_w = Inches(3.68)
    p_h = Inches(3.4)
    p_top = Inches(1.8)

    for i, (step_title, accent_col, step_bullets) in enumerate(pipeline_steps):
        left_pos = Inches(0.8 + i * 4.02)
        step_shape = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, p_top, p_w, p_h)
        step_shape.fill.solid()
        step_shape.fill.fore_color.rgb = CARD_BG
        step_shape.line.color.rgb = CARD_BORDER
        step_shape.line.width = Pt(1)
        
        stf = step_shape.text_frame
        stf.vertical_anchor = MSO_ANCHOR.TOP
        stf.word_wrap = True
        stf.margin_left = stf.margin_right = Inches(0.26)
        stf.margin_top = Inches(0.22)
        stf.margin_bottom = Inches(0.15)
        
        sp = stf.paragraphs[0]
        sp.text = step_title
        sp.font.name = FONT_FAMILY
        sp.font.size = Pt(11)
        sp.font.bold = True
        sp.font.color.rgb = accent_col
        
        for b_title, b_desc in step_bullets:
            sbp = stf.add_paragraph()
            sbp.space_before = Pt(10)
            
            r1 = sbp.add_run()
            r1.text = "• " + b_title
            r1.font.name = FONT_FAMILY
            r1.font.size = Pt(10)
            r1.font.bold = True
            r1.font.color.rgb = TEXT_MAIN
            
            r2 = sbp.add_run()
            r2.text = b_desc
            r2.font.name = FONT_FAMILY
            r2.font.size = Pt(9.5)
            r2.font.color.rgb = TEXT_MUTED

    # Bottom 2 Feature Highlight Cards
    adv_w = Inches(5.72)
    adv_h = Inches(1.55)
    adv_top = Inches(5.38)

    # Card 1: Why it is Better
    c_adv1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), adv_top, adv_w, adv_h)
    c_adv1.fill.solid()
    c_adv1.fill.fore_color.rgb = CARD_BG
    c_adv1.line.color.rgb = CARD_BORDER
    c_adv1.line.width = Pt(1)
    
    tf_adv1 = c_adv1.text_frame
    tf_adv1.vertical_anchor = MSO_ANCHOR.TOP
    tf_adv1.word_wrap = True
    tf_adv1.margin_left = tf_adv1.margin_right = Inches(0.3)
    tf_adv1.margin_top = Inches(0.18)
    
    p = tf_adv1.paragraphs[0]
    p.text = "WHY IT IS BETTER: EVIDENCE OVER KEYWORDS"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER
    
    p2 = tf_adv1.add_paragraph()
    p2.text = "Unlike traditional ATS checkers that only perform shallow string matching, ProfiQ establishes ground truth by inspecting actual repository commits, code architecture, and package manifests to verify whether a candidate can genuinely execute."
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_before = Pt(4)

    # Card 2: What is Different
    c_adv2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), adv_top, adv_w, adv_h)
    c_adv2.fill.solid()
    c_adv2.fill.fore_color.rgb = CARD_BG
    c_adv2.line.color.rgb = CARD_BORDER
    c_adv2.line.width = Pt(1)
    
    tf_adv2 = c_adv2.text_frame
    tf_adv2.vertical_anchor = MSO_ANCHOR.TOP
    tf_adv2.word_wrap = True
    tf_adv2.margin_left = tf_adv2.margin_right = Inches(0.3)
    tf_adv2.margin_top = Inches(0.18)
    
    p = tf_adv2.paragraphs[0]
    p.text = "THE INNOVATION: ACTIONABLE CAREER REMEDIATION"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_ROSE
    
    p2 = tf_adv2.add_paragraph()
    p2.text = "ProfiQ transforms evaluation from a passive 'judgment' into an active developmental roadmap. Students receive transparent explanations for every deducted point and a prescriptive milestone plan to turn red skill deficits into verified green badges."
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_before = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 4: SYSTEM ARCHITECTURE
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "03 · Architecture", "Modular Multi-Tiered System Architecture", 
               "Engineered for high-throughput profile ingestion, robust telemetry extraction, and deterministic AI scoring.")

    arch_cols = [
        ("CLIENT TIER", "React 18 + Vite", ACCENT_BLUE, [
            ("Candidate Studio: ", "Interactive profile upload, role target picker, and live readiness dials."),
            ("Visual Telemetry: ", "Radar charts, verified vs. unbacked skill matrices, and 3D score dials."),
            ("Placement Portal: ", "Cohort-wide analytics, distribution heatmaps, and batch filter tables.")
        ]),
        ("BACKEND API TIER", "Node.js + Express", ACCENT_EMERALD, [
            ("Ingestion Engine: ", "PDF parsing & multipart form handling for resumes and profile links."),
            ("Telemetry Service: ", "GitHub REST API connector extracting repo stats, languages, and commit histories."),
            ("Orchestrator: ", "Normalizes candidate profiles and manages asynchronous AI inference pipelines.")
        ]),
        ("INTELLIGENCE TIER", "Google Gemini + Engine", ACCENT_ROSE, [
            ("Skill Extraction: ", "Structured NER identifying explicit and implicit technical proficiencies."),
            ("Artifact Verifier: ", "Heuristics and LLM semantic cross-check matching code against claims."),
            ("Scoring & Roadmap: ", "Deterministic mathematical scoring and personalized 4-week sprint generator.")
        ]),
        ("PERSISTENCE TIER", "MongoDB + Schemas", ACCENT_AMBER, [
            ("Unified Profiles: ", "Stores normalized candidate claims, verified artifacts, and telemetry."),
            ("Role Benchmarks: ", "Taxonomies mapping Critical, Recommended, and Optional skills across roles."),
            ("Roadmaps & Reports: ", "Historical tracking of readiness progress over academic semesters.")
        ])
    ]

    col4_w = Inches(2.76)
    col4_h = Inches(3.65)
    col4_top = Inches(1.8)

    for i, (tier_name, tier_tech, accent_col, tier_points) in enumerate(arch_cols):
        left_pos = Inches(0.8 + i * 2.99)
        c_shape = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, col4_top, col4_w, col4_h)
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG
        c_shape.line.color.rgb = CARD_BORDER
        c_shape.line.width = Pt(1)
        
        ctf = c_shape.text_frame
        ctf.vertical_anchor = MSO_ANCHOR.TOP
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.2)
        ctf.margin_top = Inches(0.2)
        ctf.margin_bottom = Inches(0.15)
        
        cp = ctf.paragraphs[0]
        cp.text = tier_name
        cp.font.name = FONT_FAMILY
        cp.font.size = Pt(10.5)
        cp.font.bold = True
        cp.font.color.rgb = accent_col
        
        cp_sub = ctf.add_paragraph()
        cp_sub.text = tier_tech
        cp_sub.font.name = FONT_FAMILY
        cp_sub.font.size = Pt(9)
        cp_sub.font.bold = True
        cp_sub.font.color.rgb = TEXT_LIGHT
        cp_sub.space_before = Pt(2)
        
        for b_title, b_desc in tier_points:
            bp = ctf.add_paragraph()
            bp.space_before = Pt(9)
            
            r1 = bp.add_run()
            r1.text = "• " + b_title
            r1.font.name = FONT_FAMILY
            r1.font.size = Pt(9.5)
            r1.font.bold = True
            r1.font.color.rgb = TEXT_MAIN
            
            r2 = bp.add_run()
            r2.text = b_desc
            r2.font.name = FONT_FAMILY
            r2.font.size = Pt(9)
            r2.font.color.rgb = TEXT_MUTED

    eng_decisions = [
        ("KEY TECHNICAL DECISION", ACCENT_BLUE, "Modular Separation: Independent client, backend orchestrator, and AI pipelines allow replacing or upgrading LLM models without altering database or UI schemas."),
        ("HARDEST PROBLEM SOLVED", ACCENT_ROSE, "Evidence Attribution: Connecting unstructured resume claims with commit histories and repositories via semantic matching and deterministic heuristics."),
        ("VALIDATION & STABILITY", ACCENT_EMERALD, "Strict JSON Schemas: Gemini structured output constraints prevent AI hallucinations, ensuring reproducible scores and robust API contracts.")
    ]

    dec_w = Inches(3.68)
    dec_h = Inches(1.3)
    dec_top = Inches(5.62)

    for i, (dec_title, accent_col, dec_text) in enumerate(eng_decisions):
        left_pos = Inches(0.8 + i * 4.02)
        dec_shape = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, dec_top, dec_w, dec_h)
        dec_shape.fill.solid()
        dec_shape.fill.fore_color.rgb = CARD_BG
        dec_shape.line.color.rgb = CARD_BORDER
        dec_shape.line.width = Pt(1)
        
        dtf = dec_shape.text_frame
        dtf.vertical_anchor = MSO_ANCHOR.TOP
        dtf.word_wrap = True
        dtf.margin_left = dtf.margin_right = Inches(0.2)
        dtf.margin_top = Inches(0.14)
        
        dp = dtf.paragraphs[0]
        dp.text = dec_title
        dp.font.name = FONT_FAMILY
        dp.font.size = Pt(10)
        dp.font.bold = True
        dp.font.color.rgb = accent_col
        
        dp2 = dtf.add_paragraph()
        dp2.text = dec_text
        dp2.font.name = FONT_FAMILY
        dp2.font.size = Pt(9)
        dp2.font.color.rgb = TEXT_MUTED
        dp2.space_before = Pt(3)

    # -------------------------------------------------------------
    # SLIDE 5: SCORING METHODOLOGY
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "04 · Evaluation Logic", "Evidence-Based Mathematical Scoring Methodology", 
               "A transparent, explainable formula that rewards verified engineering competence and penalizes unbacked claims.")

    # Formula Hero Box
    formula_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(1.2))
    formula_box.fill.solid()
    formula_box.fill.fore_color.rgb = CARD_BG
    formula_box.line.color.rgb = ACCENT_BLUE
    formula_box.line.width = Pt(1.5)
    
    ftf = formula_box.text_frame
    ftf.vertical_anchor = MSO_ANCHOR.TOP
    ftf.word_wrap = True
    ftf.margin_left = Inches(0.4)
    ftf.margin_top = Inches(0.18)
    
    fp1 = ftf.paragraphs[0]
    fp1.text = "CORE FORMULA:   Readiness Score (R) = (S × 0.60) + (E × 0.40)"
    fp1.font.name = FONT_FAMILY
    fp1.font.size = Pt(16)
    fp1.font.bold = True
    fp1.font.color.rgb = ACCENT_BLUE
    
    fp2 = ftf.add_paragraph()
    fp2.text = "Where:  S = Skill Match Score (0–100) based on industry role benchmarks  |  E = Evidence Verification Score (0–100) based on code proof"
    fp2.font.name = FONT_FAMILY
    fp2.font.size = Pt(10.5)
    fp2.font.color.rgb = TEXT_MUTED
    fp2.space_before = Pt(4)

    # Two Main Methodology Cards
    c_s = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.2), Inches(5.72), Inches(3.72))
    c_s.fill.solid()
    c_s.fill.fore_color.rgb = CARD_BG
    c_s.line.color.rgb = CARD_BORDER
    c_s.line.width = Pt(1)
    
    stf = c_s.text_frame
    stf.vertical_anchor = MSO_ANCHOR.TOP
    stf.word_wrap = True
    stf.margin_left = stf.margin_right = Inches(0.3)
    stf.margin_top = Inches(0.24)
    
    p = stf.paragraphs[0]
    p.text = "1. SKILL MATCH SCORE (S · 60% WEIGHT)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    
    s_points = [
        ("Hierarchical Role Weights: ", "Evaluates skills against curated industry matrices according to role criticality:"),
        ("• Critical Skills (1.0× Weight): ", "Foundational technologies without which the candidate cannot operate (e.g., React & JavaScript for Frontend)."),
        ("• Recommended Skills (0.7× Weight): ", "Modern tooling and best practices that elevate technical output (e.g., TypeScript, Tailwind, REST APIs)."),
        ("• Optional Skills (0.4× Weight): ", "Value-add complementary proficiencies (e.g., Docker, GraphQL, Jest)."),
        ("Dynamic Breadth & Depth: ", "Calculates coverage ratio across core domains rather than simple keyword counting.")
    ]
    for b_title, b_desc in s_points:
        bp = stf.add_paragraph()
        bp.space_before = Pt(8)
        r1 = bp.add_run()
        r1.text = b_title
        r1.font.name = FONT_FAMILY
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = TEXT_MAIN
        r2 = bp.add_run()
        r2.text = b_desc
        r2.font.name = FONT_FAMILY
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_MUTED

    c_e = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(3.2), Inches(5.72), Inches(3.72))
    c_e.fill.solid()
    c_e.fill.fore_color.rgb = CARD_BG
    c_e.line.color.rgb = CARD_BORDER
    c_e.line.width = Pt(1)
    
    etf = c_e.text_frame
    etf.vertical_anchor = MSO_ANCHOR.TOP
    etf.word_wrap = True
    etf.margin_left = etf.margin_right = Inches(0.3)
    etf.margin_top = Inches(0.24)
    
    p = etf.paragraphs[0]
    p.text = "2. EVIDENCE VERIFICATION SCORE (E · 40% WEIGHT)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = ACCENT_ROSE
    
    e_points = [
        ("Evidence Confidence Hierarchy: ", "Each claimed skill is assigned an evidence tier backed by observable data:"),
        ("• High Confidence (0.85 – 1.0): ", "Public repository with verified commits, active commit velocity, and matching codebase."),
        ("• Moderate Confidence (0.50 – 0.70): ", "Listed academic project, course certification, or live portfolio demo URL."),
        ("• Low / Unbacked (0.10 – 0.30): ", "Resume keyword claim with zero observable codebase proof or commit history."),
        ("Penalty Mechanism: ", "Unbacked claims dilute the overall readiness score, directly discouraging resume exaggeration.")
    ]
    for b_title, b_desc in e_points:
        bp = etf.add_paragraph()
        bp.space_before = Pt(8)
        r1 = bp.add_run()
        r1.text = b_title
        r1.font.name = FONT_FAMILY
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = TEXT_MAIN
        r2 = bp.add_run()
        r2.text = b_desc
        r2.font.name = FONT_FAMILY
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 6: PRODUCT FEATURES & CAPABILITIES
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "05 · Product Capabilities", "Dual-Perspective Platform: Students & Placement Cells", 
               "Delivering immediate diagnostic value for candidate growth and powerful cohort analytics for university administrators.")

    f_cards = [
        ("CANDIDATE STUDIO: READINESS DIAL", ACCENT_BLUE, [
            ("Composite Readiness Score: ", "High-impact visual score dial reflecting current employability level (0–100)."),
            ("Domain Radar Breakdown: ", "Visualizes candidate strength across Core Languages, Frameworks, Architecture, and Tooling."),
            ("Target Benchmark Threshold: ", "Clear indicators showing whether the candidate meets Junior, Mid, or Senior entry bars.")
        ]),
        ("EVIDENCE MATRIX & GAP VISUALIZER", ACCENT_ROSE, [
            ("Verified vs. Unbacked Claims: ", "Displays an 80%+ evidence coverage target, clearly tagging proof sources."),
            ("Critical Skill Deficits: ", "Flags essential role requirements that are missing from candidate profiles."),
            ("Interactive Inspect Mode: ", "Allows candidates to see exactly which GitHub repos validated each skill.")
        ]),
        ("PERSONALIZED MILESTONE ROADMAP", ACCENT_EMERALD, [
            ("4-Week Sprint Schedule: ", "Automated sprint breakdown focusing on closing critical skill deficits step by step."),
            ("Curated Learning Links: ", "Recommends high-quality documentation, guided projects, and tutorials."),
            ("Proof of Work Deliverables: ", "Prescribes specific GitHub project milestones to generate tangible proof.")
        ]),
        ("INSTITUTIONAL PLACEMENT INTELLIGENCE", ACCENT_AMBER, [
            ("Batch Skill Heatmaps: ", "Placement cells view cohort-wide trends and common technological weaknesses."),
            ("Pre-Placement Interventions: ", "Enables organizing targeted bootcamps before recruitment drives commence."),
            ("Verified Candidate Shortlisting: ", "Allows recruiters to filter students by demonstrable code proof rather than GPA alone.")
        ])
    ]

    card2_w = Inches(5.72)
    card2_h = Inches(2.45)

    positions = [
        (Inches(0.8), Inches(1.8)),
        (Inches(6.8), Inches(1.8)),
        (Inches(0.8), Inches(4.45)),
        (Inches(6.8), Inches(4.45))
    ]

    for i, (f_title, accent_col, f_bullets) in enumerate(f_cards):
        left_pos, top_pos = positions[i]
        c_shape = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, card2_w, card2_h)
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG
        c_shape.line.color.rgb = CARD_BORDER
        c_shape.line.width = Pt(1)
        
        ctf = c_shape.text_frame
        ctf.vertical_anchor = MSO_ANCHOR.TOP
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.28)
        ctf.margin_top = Inches(0.22)
        ctf.margin_bottom = Inches(0.15)
        
        cp = ctf.paragraphs[0]
        cp.text = f_title
        cp.font.name = FONT_FAMILY
        cp.font.size = Pt(11)
        cp.font.bold = True
        cp.font.color.rgb = accent_col
        
        for b_title, b_desc in f_bullets:
            bp = ctf.add_paragraph()
            bp.space_before = Pt(8)
            
            r1 = bp.add_run()
            r1.text = "• " + b_title
            r1.font.name = FONT_FAMILY
            r1.font.size = Pt(10)
            r1.font.bold = True
            r1.font.color.rgb = TEXT_MAIN
            
            r2 = bp.add_run()
            r2.text = b_desc
            r2.font.name = FONT_FAMILY
            r2.font.size = Pt(9.5)
            r2.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 7: FEASIBILITY, MOAT & ROADMAP
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "06 · Defensibility & Horizon", "Feasibility, Competitive Moat & Growth Horizon", 
               "Low operational run-rates, clear unfair advantages over legacy tools, and a defined post-hackathon roadmap.")

    cols_s7 = [
        ("COMPETITIVE MOAT", ACCENT_ROSE, [
            ("Traditional ATS: ", "Parses resume strings blindly; easily gamed by buzzword stuffing with zero evidence verification."),
            ("LinkedIn Profiles: ", "Relies solely on unverified self-reporting and social endorsements without code audits."),
            ("Generic AI Chatbots: ", "Provides generic advice without ingesting authentic developer artifacts or scoring mathematically."),
            ("The ProfiQ Advantage: ", "Only system combining multi-platform proof of work + industry benchmark alignment + automated remediation.")
        ]),
        ("FEASIBILITY & RUN-RATE", ACCENT_AMBER, [
            ("Minimal API Costs: ", "Structured JSON prompts and token optimization keep LLM execution costs to fractions of a cent per profile."),
            ("Deterministic Caching: ", "Caches GitHub repository metadata and benchmark matrices to eliminate redundant API calls."),
            ("Zero Infrastructure Barrier: ", "Built on lightweight Node.js and MongoDB; easily deployed to serverless environments."),
            ("Privacy & Security: ", "Inspects only public code artifacts; no intrusive access to private student credentials.")
        ]),
        ("FUTURE DEVELOPMENT HORIZON", ACCENT_EMERALD, [
            ("Phase 1: Multi-Platform Ingestion: ", "Expand connectors to LeetCode, Codeforces, HackerRank, and GitLab for algorithmic telemetry."),
            ("Phase 2: University SIS Integration: ", "Automated batch sync with university ERPs to track student growth across semesters."),
            ("Phase 3: Recruiter Search Mode: ", "Allow tech recruiters to query candidate pools by verified skills (e.g., 'React + Docker > 85% proof')."),
            ("Phase 4: AI Mock Interviewer: ", "Interactive voice agent that quizzes candidates specifically on their verified code repositories.")
        ])
    ]

    for i, (col_title, accent_col, bullet_items) in enumerate(cols_s7):
        left_pos = Inches(0.8 + i * 4.02)
        c_shape = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(1.8), Inches(3.68), Inches(5.1))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG
        c_shape.line.color.rgb = CARD_BORDER
        c_shape.line.width = Pt(1)
        
        ctf = c_shape.text_frame
        ctf.vertical_anchor = MSO_ANCHOR.TOP
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.25)
        ctf.margin_top = Inches(0.25)
        ctf.margin_bottom = Inches(0.2)
        
        cp = ctf.paragraphs[0]
        cp.text = col_title
        cp.font.name = FONT_FAMILY
        cp.font.size = Pt(11)
        cp.font.bold = True
        cp.font.color.rgb = accent_col
        
        for b_title, b_desc in bullet_items:
            bp = ctf.add_paragraph()
            bp.space_before = Pt(11)
            
            r1 = bp.add_run()
            r1.text = "• " + b_title
            r1.font.name = FONT_FAMILY
            r1.font.size = Pt(10)
            r1.font.bold = True
            r1.font.color.rgb = TEXT_MAIN
            
            r2 = bp.add_run()
            r2.text = b_desc
            r2.font.name = FONT_FAMILY
            r2.font.size = Pt(9.5)
            r2.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 8: THANK YOU & CONCLUSION
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    
    # Pill
    ty_pill = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.8), Inches(0.85), Inches(3.733), Inches(0.38))
    ty_pill.fill.solid()
    ty_pill.fill.fore_color.rgb = ACCENT_ROSE_BG
    ty_pill.line.color.rgb = RGBColor(254, 205, 211)
    ty_tf = ty_pill.text_frame
    ty_tf.vertical_anchor = MSO_ANCHOR.TOP
    ty_tf.margin_left = ty_tf.margin_right = 0
    ty_tf.margin_top = Inches(0.05)
    ty_tf.margin_bottom = 0
    ty_p = ty_tf.paragraphs[0]
    ty_p.text = "DATAQUEST 3.0 · ROUND 1 · TRACK DQWL"
    ty_p.alignment = PP_ALIGN.CENTER
    ty_p.font.name = FONT_FAMILY
    ty_p.font.size = Pt(10)
    ty_p.font.bold = True
    ty_p.font.color.rgb = ACCENT_ROSE

    # Hero "THANK YOU" Title
    ty_title_box = s8.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.733), Inches(1.5))
    ty_tf2 = ty_title_box.text_frame
    ty_tf2.vertical_anchor = MSO_ANCHOR.TOP
    ty_tf2.word_wrap = True
    ty_tf2.margin_left = ty_tf2.margin_top = ty_tf2.margin_right = ty_tf2.margin_bottom = 0
    
    p = ty_tf2.paragraphs[0]
    p.text = "THANK YOU"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_FAMILY
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN
    
    p_sub = ty_tf2.add_paragraph()
    p_sub.text = "ProfiQ — Turning Resume Claims into Verifiable Career Readiness"
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.font.name = FONT_FAMILY
    p_sub.font.size = Pt(17)
    p_sub.font.bold = True
    p_sub.font.color.rgb = ACCENT_BLUE
    p_sub.space_before = Pt(6)

    # 3 Summary Cards
    ty_cards = [
        ("PROJECT VISION", ACCENT_BLUE, [
            ("From Claims to Proof: ", "Empowering engineering students to back their claims with demonstrable code."),
            ("Data-Driven Placement: ", "Giving university placement cells actionable cohort telemetry."),
            ("Hiring Confidence: ", "Providing technical recruiters with verified talent pipelines.")
        ]),
        ("TEAM [ 3Cuties ]", ACCENT_AMBER, [
            ("Aman Tyagi ", "(25BCE1710) — AI & ML Lead"),
            ("Siddhant Asati ", "(25BCE1633) — Backend & Systems Lead"),
            ("Sahaj Kumar ", "(25BCE1620) — Frontend & UX Lead"),
            ("Institution: ", "Vellore Institute of Technology (VIT)")
        ]),
        ("OPEN FOR Q&A", ACCENT_EMERALD, [
            ("Live Demonstration: ", "Ready for walkthrough of the Candidate Studio & Placement Dashboard."),
            ("Codebase & Tests: ", "Modular monorepo with automated unit & end-to-end test suites."),
            ("Discussion: ", "We welcome questions on our scoring formula and verification heuristics.")
        ])
    ]

    ty_card_w = Inches(3.68)
    ty_card_h = Inches(3.6)
    ty_top = Inches(3.2)

    for i, (c_title, accent_col, c_bullets) in enumerate(ty_cards):
        left_pos = Inches(0.8 + i * 4.02)
        c_shape = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, ty_top, ty_card_w, ty_card_h)
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG
        c_shape.line.color.rgb = CARD_BORDER
        c_shape.line.width = Pt(1)
        
        ctf = c_shape.text_frame
        ctf.vertical_anchor = MSO_ANCHOR.TOP
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.28)
        ctf.margin_top = Inches(0.25)
        ctf.margin_bottom = Inches(0.2)
        
        cp = ctf.paragraphs[0]
        cp.text = c_title
        cp.font.name = FONT_FAMILY
        cp.font.size = Pt(11.5)
        cp.font.bold = True
        cp.font.color.rgb = accent_col
        
        for b_title, b_desc in c_bullets:
            bp = ctf.add_paragraph()
            bp.space_before = Pt(10)
            
            r1 = bp.add_run()
            r1.text = "• " + b_title
            r1.font.name = FONT_FAMILY
            r1.font.size = Pt(10.5)
            r1.font.bold = True
            r1.font.color.rgb = TEXT_MAIN
            
            r2 = bp.add_run()
            r2.text = b_desc
            r2.font.name = FONT_FAMILY
            r2.font.size = Pt(10)
            r2.font.color.rgb = TEXT_MUTED

    output_path = os.path.abspath("ProfiQ_Hackathon_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation successfully saved to: {output_path}")

if __name__ == "__main__":
    create_presentation()
